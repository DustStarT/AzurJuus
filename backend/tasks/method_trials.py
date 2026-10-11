"""Declarative, deterministic acceptance for arbitrary local file/document methods."""
import hashlib
import json
from pathlib import PurePosixPath


def safe_path(value):
    if (not isinstance(value,str) or not value or len(value)>240 or '\\' in value or ':' in value
            or '\0' in value or PurePosixPath(value).is_absolute() or '..' in PurePosixPath(value).parts
            or any(p.startswith('.juus') for p in PurePosixPath(value).parts)):
        raise ValueError('试用文件必须是独立工作区内的相对路径。')
    return value


def validate_cases(cases):
    if not isinstance(cases,list) or len(cases)!=3:raise ValueError('演化需要三个不同的隔离验收场景。')
    identities=set()
    for case in cases:
        if not isinstance(case,dict):raise ValueError('试用场景必须是对象。')
        if not isinstance(case.get('prompt'),str) or not 1<=len(case['prompt'].strip())<=4000:
            raise ValueError('试用场景缺少具体要求。')
        files,expected=case.get('files'),case.get('expected')
        if not isinstance(files,dict) or len(files)>24 or not isinstance(expected,dict) or not 1<=len(expected)<=24:
            raise ValueError('每个场景必须包含 files 和 1–24 项独立 expected 验收断言。')
        for path,value in files.items():
            safe_path(path)
            if isinstance(value,str):continue
            if not isinstance(value,dict) or value.get('kind') not in {'docx','xlsx'}:
                raise ValueError('夹具仅支持文字、DOCX 或 XLSX 数据。')
            if value['kind']=='docx' and not isinstance(value.get('content'),str):raise ValueError('DOCX 夹具需要文字。')
            if value['kind']=='xlsx' and not isinstance(value.get('sheets'),dict):raise ValueError('XLSX 夹具需要 sheets 对象。')
            if value['kind']=='xlsx':validate_sheets(value['sheets'])
        for path,assertion in expected.items():
            safe_path(path)
            if not isinstance(assertion,dict) or len(assertion)!=1:raise ValueError('每项断言须指定一个验收操作。')
            kind,value=next(iter(assertion.items()))
            if kind not in {'text','contains','json','docx','xlsx','absent','sha256'}:raise ValueError('不支持的验收断言。')
            if kind in {'text','contains','docx','sha256'} and not isinstance(value,str):raise ValueError('文字断言类型无效。')
            if kind=='contains' and not value:raise ValueError('contains 断言不能为空。')
            if kind=='absent' and value is not True:raise ValueError('absent 必须为 true。')
            if kind=='sha256' and (len(value)!=64 or any(c not in '0123456789abcdef' for c in value)):
                raise ValueError('sha256 断言必须为真实摘要。')
            if kind=='xlsx' and not isinstance(value,dict):raise ValueError('XLSX 断言须指定各表的行数据。')
            if kind=='xlsx':validate_sheets(value)
        preserve=case.get('preserve',[p for p in files if p not in expected])
        if not isinstance(preserve,list) or any(p not in files for p in preserve):raise ValueError('preserve 只能指定实际夹具路径。')
        encoded=json.dumps(case,ensure_ascii=False,sort_keys=True,allow_nan=False)
        if len(encoded)>100000:raise ValueError('试用场景内容过大。')
        identities.add(hashlib.sha256(json.dumps({'files':files,'expected':expected},sort_keys=True,ensure_ascii=False).encode()).hexdigest())
    if len(identities)!=3:raise ValueError('三个试用场景的输入和验收必须不同，不能仅修改名称。')
    return cases


def validate_sheets(sheets):
    if not 1<=len(sheets)<=8:raise ValueError('XLSX 试用需要 1–8 张工作表。')
    for name,rows in sheets.items():
        if not isinstance(name,str) or not 1<=len(name)<=31 or any(c in name for c in '[]:*?/\\'):
            raise ValueError('XLSX 工作表名称无效。')
        if not isinstance(rows,list) or len(rows)>200:raise ValueError('XLSX 试用每表最多 200 行。')
        if any(not isinstance(row,list) or len(row)>50 or any(not isinstance(cell,(str,int,float,bool,type(None))) for cell in row) for row in rows):
            raise ValueError('XLSX 行必须是最多 50 个标量单元格的数组。')


def create_files(work,case):
    for name,value in case['files'].items():
        path=work/safe_path(name);path.parent.mkdir(parents=True,exist_ok=True)
        if isinstance(value,str):path.write_text(value,encoding='utf-8')
        elif value['kind']=='docx':
            from docx import Document
            document=Document()
            for line in value['content'].splitlines():document.add_paragraph(line)
            document.save(path)
        else:
            from openpyxl import Workbook
            book=Workbook();book.remove(book.active)
            for title,rows in value['sheets'].items():
                sheet=book.create_sheet(title)
                for row in rows:sheet.append(row)
            book.save(path);book.close()
    return {p:hashlib.sha256((work/p).read_bytes()).hexdigest() for p in case['files']}


def evaluate(work,case,before):
    for name,assertion in case['expected'].items():
        path=work/name
        if path.resolve().is_relative_to(work.resolve()) is False:return False
        kind,value=next(iter(assertion.items()))
        if kind=='absent':
            if path.exists():return False
            continue
        if not path.is_file():return False
        try:
            if kind=='sha256':actual=hashlib.sha256(path.read_bytes()).hexdigest()
            elif kind=='json':actual=json.loads(path.read_text(encoding='utf-8'))
            elif kind=='docx':
                from docx import Document
                actual='\n'.join(p.text for p in Document(path).paragraphs)
            elif kind=='xlsx':
                from openpyxl import load_workbook
                book=load_workbook(path,read_only=True,data_only=False)
                try:actual={s.title:[list(r) for r in s.iter_rows(values_only=True)] for s in book}
                finally:book.close()
            else:actual=path.read_text(encoding='utf-8')
            if kind=='contains':
                if value not in actual:return False
            elif not equivalent(actual,value):return False
        except Exception:return False
    for name in case.get('preserve',[p for p in case['files'] if p not in case['expected']]):
        path=work/name
        if not path.resolve().is_relative_to(work.resolve()) or not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest()!=before[name]:return False
    permitted=set(case['files'])|set(case['expected'])
    return all(p.relative_to(work).as_posix() in permitted for p in work.rglob('*') if p.is_file())


def equivalent(actual,expected):
    # Python treats False == 0 and True == 1; JSON/Excel acceptance must not.
    if isinstance(actual,bool) or isinstance(expected,bool):return type(actual) is type(expected) and actual==expected
    if isinstance(expected,dict):
        return isinstance(actual,dict) and actual.keys()==expected.keys() and all(equivalent(actual[k],v) for k,v in expected.items())
    if isinstance(expected,list):
        return isinstance(actual,list) and len(actual)==len(expected) and all(equivalent(a,b) for a,b in zip(actual,expected))
    return actual==expected
