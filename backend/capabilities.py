"""Task-scoped local tools. Browser and desktop share the same audited gateway."""
from __future__ import annotations

import asyncio
import base64
import hashlib
import json
import os
import shutil
import sys
import time
from pathlib import Path
from uuid import uuid4


CATALOG = {
    "discuss": "Ask or challenge an assigned teammate: actorId, text (1-2000 chars), optional kind (question/suggestion/objection/answer/resolved), threadId to continue a discussion. Maximum two exchanges per thread, 8 per task. Replies appear in group chat and tool results. To mark resolved supply threadId and explanation. Discussion is not execution evidence.",
    "list_dir": "List path (default .).",
    "read_file": "Read UTF-8 text by path, offset (line, default 0), column (default 0), limit (lines, default 200), maxChars (default 20000). Continue at nextOffset/nextColumn.",
    "search": "Search path recursively for literal pattern; returns paths, line numbers and text.",
    "write_file": "Write UTF-8 content to path; existing files are snapshotted.",
    "patch_file": "Replace exactly one occurrence of old with new in path; fails if ambiguous.",
    "mkdir": "Create path within the workspace.",
    "copy": "Copy src file to dest; overwrite requires approval.",
    "move": "Move src file to dest; overwrite requires approval.",
    "restore": "Restore path from snapshotId; requires approval.",
    "undo": "Undo a copy/move using its snapshotId; verifies destination has not changed and restores overwritten destination. Requires approval.",
    "read_document": "Read PDF, DOCX or XLSX path; offset and limit page/paragraph/row units. Includes source locations.",
    "write_document": "Write DOCX from content, or XLSX from sheets mapping sheet names to lists of rows.",
    "command": "Run argv (array, no shell interpolation) in cwd within workspace. timeout seconds default 120. Requires task command grant.",
    "browser": "Browser op: navigate(url), read, click(selector), fill(selector,value), submit(selector), download(selector,path), screenshot, tabs, select(index). Selectors are Playwright selectors. Submission/click requires approval.",
    "desktop": "Windows op: windows, inspect(handle), focus(handle), select(handle,index), open(handle,index), invoke(handle,index), type(handle,index,text), shortcut(handle,action: copy/cut/paste/undo/rename/select_all/up), explorer(path), screenshot, click(x,y). Element indices and shortcuts require a recent inspect. UI actions require approval.",
    "deliver": "Submit result: summary, artifacts [paths], checks [{callId,description}], unresolved [strings], optional notes [strings]. Read back created files. Read-only answers/listings may use artifacts=[] and successful list_dir/read calls as checks. unresolved ONLY means unmet explicit user requirements; optional further analysis and scope limitations go in notes/summary. Resolve prior failed/rejected operations using resolvedErrors [{callId,resolution,checkCallId}] with later verification, or list genuine unresolved issues.",
    "delegate": "Propose host-scheduled assignments: tasks [{id,actorId,brief,dependsOn:[id],acceptance:[strings]}]. Only secretary planner may use this.",
    "execution_history": "Read this task's durable tool journal. Supply callId for one call, or offset/limit for call metadata. Full call JSON is paginated by charOffset/maxChars; nextCharOffset locates the remainder.",
}
MUTATING = {"write_file", "patch_file", "mkdir", "copy", "move", "restore", "undo", "write_document"}


def phase_actions(phase):
    if phase == "chat" or phase.startswith("discussion_"):
        return []
    if phase == "planner":
        return ["delegate", "execution_history"]
    if phase == "reviewer":
        return ["read_file", "read_document", "search", "list_dir", "deliver", "execution_history"]
    return [name for name in CATALOG if name != "delegate"]


class CapabilityError(ValueError):
    pass


class LocalCapabilities:
    def __init__(self, data_root: Path):
        self.data_root = data_root
        self.write_lock = asyncio.Lock()
        self.desktop_lock = asyncio.Lock()
        self.browsers = {}
        self.processes = {}
        self._desktop_elements = {}
        self.desktop_owner = None
        self.browser_locks = {}
        self._screenshots = {}

    def release_desktop(self, run_id):
        if self.desktop_owner == run_id:
            self.desktop_owner = None
        for key in list(self._desktop_elements):
            if key[0] == run_id:
                self._desktop_elements.pop(key, None)
        self._screenshots.pop(run_id, None)

    async def release_browser(self, run_id):
        browser = self.browsers.pop(run_id, None)
        if browser:
            pw, context, _ = browser
            try:
                await context.close()
            finally:
                await pw.stop()
        self.browser_locks.pop(run_id, None)

    async def threaded(self, function, *args):
        # Cancellation cannot stop a native thread. Retain the lease until it exits.
        task = asyncio.create_task(asyncio.to_thread(function, *args))
        try:
            return await asyncio.shield(task)
        except asyncio.CancelledError:
            try:
                await task
            finally:
                raise

    def path(self, root: str, value=".") -> Path:
        base = Path(root).resolve(strict=True)
        candidate = (base / str(value)).resolve()
        if not candidate.is_relative_to(base):
            raise CapabilityError("路径超出授权工作区。")
        if candidate.is_relative_to(self.data_root.resolve()):
            raise CapabilityError("运行时状态和凭据目录不能作为任务文件操作目标。")
        return candidate

    def approval_reason(self, run, name, args):
        if name in {"restore", "undo"}:
            return "恢复快照将覆盖当前文件"
        if name in {"copy", "move"} and self.path(run["workspace"], args["dest"]).exists():
            return "目标文件已存在，操作将覆盖内容"
        if name == "command" and "command" not in run.get("grants", []):
            return "允许此任务执行本地程序；程序具有当前用户权限，目录限制不构成进程沙箱"
        if name == "command":
            cmd = " ".join(args.get("argv", [])).lower()
            if any(s in cmd for s in ("remove-item", "rmdir", "rm ", "del ", "format", "shutdown", "reg.exe", "curl ", "invoke-webrequest", "push", "upload")):
                return "命令可能删除数据、更改系统或向外部提交"
        if name == "browser" and args.get("op") in {"click", "submit", "download"}:
            return "网页点击或提交可能向外部发送数据，请确认目标与内容"
        if name == "desktop" and args.get("op") in {"invoke", "type", "click", "select", "open", "shortcut"}:
            return "确认在目标 Windows 窗口执行操作"
        return None

    def snapshot(self, path: Path, root: str):
        if not path.is_file():
            return None
        sid = uuid4().hex
        folder = self.data_root / "snapshots"
        folder.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, folder / sid)
        (folder / (sid + ".json")).write_text(json.dumps({"path": str(path), "workspace": root}), encoding="utf-8")
        return sid

    async def execute(self, run, name, args, call_id, emit):
        if name not in CATALOG:
            raise CapabilityError("未知工具。")
        if name == "command":
            async with self.write_lock:
                return await self.command(run, args, call_id, emit)
        if name == "browser":
            async with self.browser_locks.setdefault(run["id"], asyncio.Lock()):
                return await self.browser(run, args)
        if name == "desktop":
            if self.desktop_owner not in {None, run["id"]}:
                raise CapabilityError("桌面正在由另一个任务独占，请稍后继续。")
            self.desktop_owner = run["id"]
            async with self.desktop_lock:
                async with self.write_lock:
                    return await self.threaded(self.desktop, run, args)
        if name in MUTATING:
            async with self.write_lock:
                return await self.threaded(self.file_tool, run, name, args)
        return await self.threaded(self.file_tool, run, name, args)

    def file_tool(self, run, name, args):
        root = run["workspace"]
        path = self.path(root, args.get("path", "."))
        offset = max(0, int(args.get("offset", 0)))
        limit = max(1, min(500, int(args.get("limit", 200))))
        if name == "list_dir":
            entries = sorted(path.iterdir(), key=lambda p: (p.is_file(), p.name.lower()))
            return {"entries": [{"name": p.name, "directory": p.is_dir()} for p in entries[offset:offset + limit]], "total": len(entries), "nextOffset": offset + limit if offset + limit < len(entries) else None}
        if name == "read_file":
            raw = path.read_bytes()
            lines = raw.decode("utf-8-sig").splitlines()
            column = max(0, int(args.get("column", 0)))
            budget = max(100, min(50000, int(args.get("maxChars", 20000))))
            content, used, next_line, next_column = [], 0, offset, column
            for i in range(offset, min(len(lines), offset + limit)):
                start = column if i == offset else 0
                prefix = f"{i + 1}: "
                capacity = budget - used - len(prefix) - 1
                if capacity <= 0:
                    break
                chunk = lines[i][start:start + capacity]
                content.append(prefix + chunk)
                used += len(prefix) + len(chunk) + 1
                next_line, next_column = (i, start + len(chunk)) if start + len(chunk) < len(lines[i]) else (i + 1, 0)
                if next_column:
                    break
            return {"path": str(path.relative_to(Path(root).resolve())), "content": "\n".join(content), "bytes": len(raw), "totalLines": len(lines), "nextOffset": next_line if next_line < len(lines) else None, "nextColumn": next_column, "sha256": hashlib.sha256(raw).hexdigest()}
        if name == "search":
            matches = []
            pattern = str(args["pattern"])
            for base, dirs, files in os.walk(path, followlinks=False):
                dirs[:] = [d for d in dirs if d not in {".git", ".venv", "node_modules", ".azurjuus", ".vendor"} and not (Path(base) / d).is_symlink() and not (hasattr(Path(base) / d, "is_junction") and (Path(base) / d).is_junction())]
                for filename in files:
                    p = self.path(root, str(Path(base) / filename))
                    if p.stat().st_size > 5_000_000:
                        continue
                    try:
                        for number, line in enumerate(p.read_text(encoding="utf-8").splitlines(), 1):
                            if pattern in line:
                                matches.append({"path": str(p.relative_to(Path(root).resolve())), "line": number, "text": line[:1000]})
                    except (UnicodeError, OSError):
                        continue
                    if len(matches) >= offset + limit + 1:
                        return {"matches": matches[offset:offset + limit], "nextOffset": offset + limit}
            return {"matches": matches[offset:offset + limit], "nextOffset": None}
        if name == "mkdir":
            path.mkdir(parents=True, exist_ok=True)
            return {"path": str(path), "exists": path.is_dir()}
        if name in {"write_file", "patch_file"}:
            content = str(args.get("content", ""))
            if name == "patch_file":
                content = path.read_text(encoding="utf-8")
                if not args.get("old") or content.count(args["old"]) != 1:
                    raise CapabilityError("补丁原文必须且只能匹配一次。")
                content = content.replace(args["old"], args["new"], 1)
            sid = self.snapshot(path, root)
            path.parent.mkdir(parents=True, exist_ok=True)
            temporary = path.with_name(path.name + ".azur-" + uuid4().hex)
            temporary.write_text(content, encoding="utf-8")
            temporary.replace(path)
            return {"path": str(path), "snapshotId": sid, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "bytes": path.stat().st_size}
        if name in {"copy", "move"}:
            src, dest = self.path(root, args["src"]), self.path(root, args["dest"])
            if not src.is_file() or src == dest:
                raise CapabilityError("源文件不存在或源与目标相同。")
            sid, dest_sid = self.snapshot(src, root), self.snapshot(dest, root)
            dest.parent.mkdir(parents=True, exist_ok=True)
            (shutil.copy2 if name == "copy" else shutil.move)(str(src), str(dest))
            meta_path = self.data_root / "snapshots" / (sid + ".json")
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            meta["undo"] = {"operation": name, "destination": str(dest), "destinationSnapshotId": dest_sid, "sha256": hashlib.sha256(dest.read_bytes()).hexdigest()}
            meta_path.write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
            return {"path": str(dest), "snapshotId": sid, "destinationSnapshotId": dest_sid, "sha256": hashlib.sha256(dest.read_bytes()).hexdigest()}
        if name == "undo":
            sid = str(args["snapshotId"])
            if not sid.isalnum():
                raise CapabilityError("无效快照。")
            folder = self.data_root / "snapshots"
            meta_path = folder / (sid + ".json")
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            operation = meta.get("undo")
            if not operation or operation.get("done"):
                raise CapabilityError("此快照不包含可撤销操作，或已撤销。")
            src = self.path(root, meta["path"])
            dest = self.path(root, operation["destination"])
            if not dest.is_file() or hashlib.sha256(dest.read_bytes()).hexdigest() != operation["sha256"]:
                raise CapabilityError("目标文件已改变，不能自动撤销。")
            if operation["operation"] == "move" and src.exists():
                raise CapabilityError("源路径已有新文件，不能自动撤销。")
            destination_snapshot = operation.get("destinationSnapshotId")
            if destination_snapshot:
                destination_meta = json.loads((folder / (destination_snapshot + ".json")).read_text(encoding="utf-8"))
                if self.path(root, destination_meta["path"]) != dest:
                    raise CapabilityError("目标快照不匹配。")
            self.snapshot(dest, root)
            if operation["operation"] == "move":
                src.parent.mkdir(parents=True, exist_ok=True)
                shutil.move(str(dest), str(src))
            else:
                dest.unlink()
            if destination_snapshot:
                shutil.copy2(folder / destination_snapshot, dest)
            operation["done"] = True
            meta_path.write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
            return {"undone": True, "source": str(src), "destinationRestored": bool(destination_snapshot)}
        if name == "restore":
            sid = str(args["snapshotId"])
            if not sid.isalnum():
                raise CapabilityError("无效快照。")
            folder = self.data_root / "snapshots"
            meta = json.loads((folder / (sid + ".json")).read_text(encoding="utf-8"))
            if self.path(root, meta["path"]) != path:
                raise CapabilityError("快照与目标路径不匹配。")
            current = self.snapshot(path, root)
            path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(folder / sid, path)
            return {"path": str(path), "snapshotId": current, "restored": True}
        if name == "read_document":
            suffix = path.suffix.lower()
            items = []
            if suffix == ".pdf":
                from pypdf import PdfReader
                reader = PdfReader(path)
                items = [{"page": i + 1, "text": p.extract_text() or "", "needsOCR": not bool(p.extract_text())} for i, p in enumerate(reader.pages)]
            elif suffix == ".docx":
                from docx import Document
                doc = Document(path)
                items = [{"paragraph": i + 1, "text": p.text} for i, p in enumerate(doc.paragraphs)]
                items.extend({"table": i + 1, "rows": [[c.text for c in r.cells] for r in t.rows]} for i, t in enumerate(doc.tables))
            elif suffix == ".xlsx":
                from openpyxl import load_workbook
                book = load_workbook(path, read_only=True, data_only=False)
                try:
                    items = [{"sheet": sheet.title, "row": i + 1, "cells": list(row)} for sheet in book for i, row in enumerate(sheet.iter_rows(values_only=True))]
                finally:
                    book.close()
            else:
                raise CapabilityError("支持 PDF、DOCX、XLSX。")
            return {"path": str(path), "items": items[offset:offset + limit], "total": len(items), "nextOffset": offset + limit if offset + limit < len(items) else None, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        if name == "write_document":
            sid = self.snapshot(path, root)
            path.parent.mkdir(parents=True, exist_ok=True)
            if path.suffix.lower() == ".docx":
                from docx import Document
                doc = Document()
                for line in args.get("content", "").splitlines():
                    doc.add_paragraph(line)
                doc.save(path)
            elif path.suffix.lower() == ".xlsx":
                from openpyxl import Workbook
                book = Workbook()
                book.remove(book.active)
                for title, rows in args["sheets"].items():
                    sheet = book.create_sheet(title)
                    for row in rows:
                        sheet.append(row)
                book.save(path)
                book.close()
            else:
                raise CapabilityError("输出格式必须为 DOCX 或 XLSX。")
            return {"path": str(path), "snapshotId": sid, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        raise CapabilityError("该工具由任务调度器处理。")

    async def command(self, run, args, call_id, emit):
        argv = args.get("argv")
        if not isinstance(argv, list) or not argv or not all(isinstance(v, str) and "\0" not in v for v in argv):
            raise CapabilityError("argv 必须为非空字符串数组。")
        cwd = self.path(run["workspace"], args.get("cwd", "."))
        env = {k: v for k, v in os.environ.items() if not any(s in k.upper() for s in ("TOKEN", "SECRET", "PASSWORD", "API_KEY"))}
        process = await asyncio.create_subprocess_exec(*argv, cwd=cwd, env=env, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.STDOUT, creationflags=0x08000000 if os.name == "nt" else 0)
        self.processes[call_id] = process
        output = []
        log_dir = self.data_root / "command-logs"
        log_dir.mkdir(parents=True, exist_ok=True)
        try:
            async with asyncio.timeout(max(1, min(1800, int(args.get("timeout", 120))))):
                with (log_dir / (call_id + ".log")).open("wb") as log:
                    while chunk := await process.stdout.read(4096):
                        log.write(chunk)
                        text = chunk.decode("utf-8", errors="replace")
                        output.append(text)
                        if sum(map(len, output)) > 32000:
                            output = output[-4:]
                        emit("tool.output", {"callId": call_id, "text": text})
                    await process.wait()
            return {"exitCode": process.returncode, "output": "".join(output), "logId": call_id, "success": process.returncode == 0}
        except TimeoutError as exc:
            await self.kill(process)
            raise TimeoutError("命令执行超时，进程已停止。") from exc
        except asyncio.CancelledError:
            await self.kill(process)
            raise
        finally:
            self.processes.pop(call_id, None)

    async def kill(self, process):
        if process.returncode is not None:
            return
        if os.name == "nt":
            killer = await asyncio.create_subprocess_exec("taskkill", "/PID", str(process.pid), "/T", "/F", stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL, creationflags=0x08000000)
            await killer.wait()
        else:
            process.kill()
        await process.wait()

    async def browser(self, run, args):
        rid = run["id"]
        if rid not in self.browsers:
            from playwright.async_api import async_playwright
            pw = await async_playwright().start()
            try:
                context = await pw.chromium.launch_persistent_context(str(self.data_root / "browsers" / rid), headless=bool(run.get("headless", False)), accept_downloads=True)
            except BaseException:
                await pw.stop()
                raise
            self.browsers[rid] = (pw, context, context.pages[0] if context.pages else await context.new_page())
        pw, context, page = self.browsers[rid]
        op = args["op"]
        if op == "navigate":
            from urllib.parse import urlparse
            if urlparse(args["url"]).scheme not in {"http", "https"}:
                raise CapabilityError("浏览器只允许 HTTP(S) 页面。")
            await page.goto(args["url"], wait_until="domcontentloaded", timeout=30000)
        elif op == "fill":
            await page.locator(args["selector"]).fill(args["value"], timeout=10000)
        elif op in {"click", "submit"}:
            await page.locator(args["selector"]).click(timeout=10000)
        elif op == "download":
            path = self.path(run["workspace"], args["path"])
            if path.exists():
                raise CapabilityError("下载目标已存在，请使用新文件名。")
            async with page.expect_download() as download:
                await page.locator(args["selector"]).click(timeout=10000)
            async with self.write_lock:
                path.parent.mkdir(parents=True, exist_ok=True)
                if path.exists():
                    raise CapabilityError("下载目标在等待期间已被创建，请使用新文件名。")
                await (await download.value).save_as(path)
            return {"path": str(path), "bytes": path.stat().st_size}
        elif op == "tabs":
            return {"tabs": [{"index": i, "url": p.url} for i, p in enumerate(context.pages)]}
        elif op == "select":
            page = context.pages[int(args["index"])]
            self.browsers[rid] = pw, context, page
        elif op == "screenshot":
            if not run.get("visionEnabled"):
                raise CapabilityError("请先在设置中启用支持图像的模型。")
            return {"image": base64.b64encode(await page.screenshot()).decode(), "mimeType": "image/png"}
        elif op != "read":
            raise CapabilityError("未知浏览器操作。")
        content = await page.locator("body").inner_text()
        offset = max(0, int(args.get("offset", 0)))
        return {"url": page.url, "title": await page.title(), "content": content[offset:offset + 20000], "nextOffset": offset + 20000 if offset + 20000 < len(content) else None, "controls": await page.locator("a,button,input,textarea,select").evaluate_all("els => els.slice(0,100).map(e => ({tag:e.tagName,id:e.id,name:e.name,text:e.innerText?.slice(0,120),type:e.type}))")}

    def desktop(self, run, args):
        if os.name != "nt":
            raise CapabilityError("桌面工具需要 Windows。")
        import ctypes
        import pythoncom
        pythoncom.CoInitialize()
        prior_dpi = None
        try:
            user32 = ctypes.windll.user32
            from ctypes import wintypes
            if hasattr(user32, "SetThreadDpiAwarenessContext"):
                user32.SetThreadDpiAwarenessContext.argtypes = [ctypes.c_void_p]
                user32.SetThreadDpiAwarenessContext.restype = ctypes.c_void_p
                prior_dpi = user32.SetThreadDpiAwarenessContext(ctypes.c_void_p(-4))
            user32.OpenInputDesktop.restype = wintypes.HANDLE
            user32.OpenInputDesktop.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
            user32.CloseDesktop.argtypes = [wintypes.HANDLE]
            user32.GetForegroundWindow.restype = wintypes.HWND
            handle = user32.OpenInputDesktop(0, False, 0x0100)
            if not handle:
                raise CapabilityError("Windows 桌面已锁定或无法访问。")
            user32.CloseDesktop(handle)
            from pywinauto import Desktop
            desktop = Desktop(backend="uia")
            op = args["op"]
            if op == "windows":
                return {"windows": [{"handle": w.handle, "title": w.window_text()} for w in desktop.windows() if w.window_text()]}
            if op == "explorer":
                import subprocess
                path = self.path(run["workspace"], args.get("path", "."))
                subprocess.Popen(["explorer.exe", "/n,", str(path)])
                return {"opened": str(path)}
            if op == "screenshot":
                if not run.get("visionEnabled"):
                    raise CapabilityError("桌面截图定位需要启用视觉模型。")
                from PIL import ImageGrab
                import io
                data = io.BytesIO()
                ImageGrab.grab(all_screens=True).save(data, format="PNG")
                origin = (user32.GetSystemMetrics(76), user32.GetSystemMetrics(77))
                self._screenshots[run["id"]] = (time.monotonic(), user32.GetForegroundWindow(), origin)
                return {"image": base64.b64encode(data.getvalue()).decode(), "mimeType": "image/png", "screenOrigin": {"x":origin[0],"y":origin[1]}, "coordinateSystem": "physical screen pixels; add screenOrigin to image coordinates"}
            if op == "click":
                if not run.get("visionEnabled"):
                    raise CapabilityError("坐标操作需要视觉模型与最新截图。")
                at, focused, _ = self._screenshots.get(run["id"], (0, None, (0,0)))
                if time.monotonic() - at > 30 or focused != user32.GetForegroundWindow():
                    raise CapabilityError("截图已过期或窗口焦点改变，请重新截图定位。")
                from pywinauto import mouse
                mouse.click(coords=(int(args["x"]), int(args["y"])))
                return {"clicked": True}
            window = desktop.window(handle=int(args["handle"]))
            if op == "inspect":
                elements = window.descendants()
                self._desktop_elements[(run["id"], args["handle"])] = (time.monotonic(), elements)
                return {"title": window.window_text(), "elements": [{"index": i, "name": e.window_text(), "type": e.element_info.control_type, "automationId": e.element_info.automation_id} for i, e in enumerate(elements[:250])]}
            window.set_focus()
            for _ in range(10):
                if user32.GetForegroundWindow() == int(args["handle"]):
                    break
                time.sleep(.05)
            if user32.GetForegroundWindow() != int(args["handle"]):
                raise CapabilityError(f"目标窗口未取得焦点，已停止操作。（目标 {args['handle']}，当前 {user32.GetForegroundWindow()}）")
            if op == "focus":
                return {"focused": window.window_text()}
            at, elements = self._desktop_elements.get((run["id"], args["handle"]), (0, []))
            if time.monotonic() - at > 30:
                raise CapabilityError("控件快照已过期，请重新 inspect。")
            if op == "shortcut":
                shortcuts = {"copy": "^c", "cut": "^x", "paste": "^v", "undo": "^z", "rename": "{F2}", "select_all": "^a", "up": "%{UP}"}
                action = args.get("action")
                if action not in shortcuts:
                    raise CapabilityError("未知桌面快捷操作。")
                from pywinauto.keyboard import send_keys
                send_keys(shortcuts[action], pause=.05)
                self._desktop_elements.pop((run["id"], args["handle"]), None)
                return {"success": True, "action": action, "title": window.window_text(), "note": "请重新检查窗口和文件结果。"}
            index = int(args["index"])
            if not 0 <= index < min(len(elements), 250):
                raise CapabilityError("无效控件编号。")
            element = elements[index]
            if op == "invoke":
                element.invoke()
            elif op == "select":
                element.set_focus()
                element.select()
            elif op == "open":
                element.double_click_input()
            elif op == "type":
                element.set_edit_text(str(args["text"]))
            else:
                raise CapabilityError("未知桌面操作。")
            return {"success": True, "title": window.window_text()}
        finally:
            if prior_dpi:
                user32.SetThreadDpiAwarenessContext(prior_dpi)
            pythoncom.CoUninitialize()

    async def close(self):
        for process in list(self.processes.values()):
            await self.kill(process)
        for run_id in list(self.browsers):
            await self.release_browser(run_id)
