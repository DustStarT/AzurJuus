"""Opt-in cloud acceptance: uses saved credentials, only synthetic isolated files."""
import asyncio
import argparse
import json
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import sqlite3
import sys
import threading
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from backend.credentials import reveal
from backend.tasks.run_coordinator import RunCoordinator
from backend.tasks.run_store import RunStore


async def main(collaborative=False, chat=False):
    with sqlite3.connect((ROOT/".azurjuus/azurjuus.db").as_uri()+"?mode=ro",uri=True) as db:
        row = db.execute("SELECT llm_base_url,llm_model,llm_api_key FROM workspace_settings LIMIT 1").fetchone()
    if not row or not row[2]:
        raise SystemExit("No saved model credential; cloud acceptance cannot run.")
    key = reveal(row[2])
    work = ROOT/"validation"/"cloud-work"/uuid4().hex
    work.mkdir(parents=True)
    (work/"source.txt").write_text("Synthetic acceptance data: apples=3, oranges=4. Total should be 7.\n",encoding="utf-8")
    if collaborative:
        from docx import Document
        from openpyxl import Workbook
        doc = Document()
        doc.add_paragraph("Synthetic document input: pears=5.")
        doc.save(work / "source.docx")
        book = Workbook()
        book.active.append(["fruit", "quantity"])
        book.active.append(["bananas", 6])
        book.save(work / "source.xlsx")
    settings = {"llmBaseUrl":row[0],"llmModel":row[1],"llmApiKey":key,"authorizedWorkspaceRoot":str(work),"maxTurns":30,"runTimeoutSeconds":180}
    store = RunStore(ROOT/".azurjuus"/"cloud-acceptance"/"runs.db")
    async def finished(*args): pass
    c = RunCoordinator(store,lambda:settings,finished,"")
    loop = asyncio.get_running_loop()
    class Handler(BaseHTTPRequestHandler):
        def log_message(self,*args): pass
        def do_POST(self):
            body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            token = self.headers.get("Authorization","").removeprefix("Bearer ")
            try:
                result = asyncio.run_coroutine_threadsafe(c.tool(token,body),loop).result(timeout=190)
                self.send_response(200)
            except Exception as exc:
                result = {"error":str(exc).replace(key,"[redacted]")}
                self.send_response(403 if isinstance(exc, PermissionError) else 400 if isinstance(exc, ValueError) else 500)
            self.send_header("Content-Type","application/json")
            self.end_headers()
            self.wfile.write(json.dumps(result,ensure_ascii=False).encode())
    server = ThreadingHTTPServer(("127.0.0.1",0),Handler)
    thread = threading.Thread(target=server.serve_forever,daemon=True)
    thread.start()
    c.endpoint=f"http://127.0.0.1:{server.server_port}/tool"
    prompt = "读取 source.txt，将两种水果数量相加，把依据和总数写入 report.txt。读取 report.txt 核验后调用 deliver 交付。只需要这项小任务，不执行命令，不联网，不创建其他文件。"
    actors = [{"id":"tester","name":"测试秘书","systemPrompt":"准确完成文件任务，以工具证据为准。"}]
    if chat:
        prompt = "这是隔离的闲聊连接验收。请只回复 AZUR_CHAT_OK。"
    if collaborative:
        prompt = "用恰好两个有依赖的子任务完成合成文档测试。第一项由测试分析员读取 source.txt、source.docx 和 source.xlsx 的水果数量，计算总数，将各来源、各数量和总数写入 totals.json。第二项由测试秘书读取第一项实际产物 totals.json，将来源和结论写入 report.txt。两项均需读回核验后调用 deliver。最终秘书独立核验两个文件。只操作工作区这三份输入和两个指定输出，不执行命令、不联网。"
        actors.append({"id":"analyst","name":"测试分析员","systemPrompt":"逐项读取原始资料，保留来源，准确计算，不编造。"})
    run = await c.admit({"prompt":prompt,"conversationId":"acceptance","mode":"chat" if chat else "task","collaborative":collaborative,"requestId":uuid4().hex},actors,settings,[])
    task=c.tasks[run["id"]]
    try:
        previous = None
        while not task.done():
            await asyncio.sleep(5)
            current=store.get(run["id"])
            # This fixture explicitly prohibits commands or external actions.
            # Reject an unexpected request and let the executor revise its plan.
            for pending in store.calls(run["id"]):
                if pending["status"] == "waiting_approval":
                    await c.resolve(run["id"], pending["id"], "rejected")
            progress = {"status":current["status"],"calls":len(store.calls(run["id"]))}
            if progress != previous:
                print(json.dumps(progress),flush=True)
                previous = progress
        await task
        current=store.get(run["id"])
        report={"model":row[1],"status":current["status"],"error":str(current.get("error") or "").replace(key,"[redacted]"),"artifacts":current["artifacts"],"workDirectory":str(work),"assignments":[{"id":a["id"],"actorId":a["actorId"],"dependsOn":a["dependsOn"],"status":a["status"]} for a in current["assignments"]],"calls":[{"name":v["name"],"phase":v.get("phase"),"status":v["status"]} for v in store.calls(run["id"])]}
        content = (work / "report.txt").read_text(encoding="utf-8") if (work / "report.txt").is_file() else ""
        report["outputChecks"] = {"reportExists": bool(content), "expectedTotalPresent": str(18 if collaborative else 7) in content}
        if chat:
            report["outputChecks"] = {"expectedReply": "AZUR_CHAT_OK" in str((current.get("result") or {}).get("summary", "")), "noToolCalls": not store.calls(run["id"])}
        if collaborative:
            report["outputChecks"]["dependencyHandoff"] = len(current["assignments"]) == 2 and any(a["dependsOn"] for a in current["assignments"])
            report["outputChecks"]["totalsJsonExists"] = (work / "totals.json").is_file()
        report["passed"] = current["status"] == "completed" and all(report["outputChecks"].values())
        name = "cloud-chat-report.json" if chat else "cloud-collaboration-report.json" if collaborative else "cloud-report.json"
        (ROOT/"validation"/name).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
        print(json.dumps(report,ensure_ascii=False),flush=True)
        return 0 if report["passed"] else 1
    finally:
        await c.close()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    modes = parser.add_mutually_exclusive_group()
    modes.add_argument("--collaborative", action="store_true")
    modes.add_argument("--chat", action="store_true")
    args = parser.parse_args()
    raise SystemExit(asyncio.run(main(args.collaborative, args.chat)))
